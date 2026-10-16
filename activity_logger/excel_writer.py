"""Writes events to a daily Excel workbook (one file per day).

Runs on its own thread, draining a queue and appending rows. Each day gets its
own file: logs/activity_YYYY-MM-DD.xlsx. The workbook is reopened and saved per
batch, which is cheap because we only emit an event when something changes.
"""

import os
import queue
import threading
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from . import config
from .events import Event

HEADERS = ["Date", "Time", "Event", "Application", "Detail", "Path", "Duration (s)"]


def _path_for(day: date) -> str:
    return os.path.join(config.log_dir(), f"activity_{day.isoformat()}.xlsx")


class ExcelWriter(threading.Thread):
    def __init__(self, event_queue: "queue.Queue[Event]"):
        super().__init__(name="ExcelWriter", daemon=True)
        self._queue = event_queue
        self._stop = threading.Event()

    def stop(self) -> None:
        self._stop.set()
        # Wake the queue.get() so the thread can notice the stop flag promptly.
        self._queue.put(None)  # type: ignore[arg-type]

    def run(self) -> None:
        while not self._stop.is_set():
            batch = self._drain()
            if batch:
                self._write(batch)

    def _drain(self) -> list[Event]:
        batch: list[Event] = []
        try:
            first = self._queue.get(timeout=1.0)
        except queue.Empty:
            return batch
        if first is not None:
            batch.append(first)
        # Grab anything else already waiting so we save once per burst.
        while True:
            try:
                item = self._queue.get_nowait()
            except queue.Empty:
                break
            if item is not None:
                batch.append(item)
        return batch

    def _write(self, batch: list[Event]) -> None:
        # Events in a batch almost always share one day, but split just in case
        # the clock rolls over midnight mid-batch.
        by_day: dict[date, list[Event]] = {}
        for ev in batch:
            by_day.setdefault(ev.time.date(), []).append(ev)

        for day, events in by_day.items():
            path = _path_for(day)
            try:
                if os.path.exists(path):
                    wb = load_workbook(path)
                    ws = wb.active
                else:
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "activity"
                    ws.append(HEADERS)
                    for cell in ws[1]:
                        cell.font = Font(bold=True)

                for ev in events:
                    ws.append([
                        ev.time.strftime("%Y-%m-%d"),
                        ev.time.strftime("%H:%M:%S"),
                        ev.kind,
                        ev.app,
                        ev.detail,
                        ev.path,
                        round(ev.duration, 1) if ev.duration is not None else "",
                    ])

                wb.save(path)
            except Exception as exc:  # pragma: no cover - keep logger alive
                # Never let a transient file error (e.g. file open in Excel)
                # kill the writer thread; drop the batch and carry on.
                _warn(f"could not write {path}: {exc}")


def _warn(msg: str) -> None:
    try:
        with open(os.path.join(config.log_dir(), "_logger_errors.txt"), "a",
                  encoding="utf-8") as fh:
            fh.write(f"{datetime.now().isoformat()} {msg}\n")
    except Exception:
        pass
